import os, sys, json, datetime, threading, urllib.request, urllib.parse
from http.server import HTTPServer, SimpleHTTPRequestHandler
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

EXCEL_FILE = "leads.xlsx"
CSV_FILE = "leads.csv"
JSON_FILE = "leads.json"
lead_lock = threading.Lock()

ALEM_API_URL = "https://llm.alem.ai/v1/chat/completions"
ALEM_API_KEY = "sk-I8agzhli09Od5WbFynXkyA"

FARAB_SYSTEM_PROMPT = """Сен Farab Tour (Мекке мен Мәдинаға рухани Ұмра сапарларын кәсіби деңгейде ұйымдастырушы туроператор) компаниясының ресми ақылды AI кеңесшісісің.

Компанияның негізгі мәліметтері:
- Негізін қалаушы әрі жетекші ұстаз: дінтанушы Әлфараби Сағымбекұлы (6+ жыл тәжірибе, 1200+ риза қажы, 35+ сәтті ұйымдастырылған топ).
- Ұшу қаласы: Шымкент халықаралық әуежайынан тікелей чартерлік/тұрақты рейстермен Мәдина немесе Жиддаға екі жаққа. Сондай-ақ Алматы мен Астанадан қосылу мүмкіндігі бар.
- Негізгі хит пакет: DOSTYK PACKAGE (11 күндік толық сапар — Мәдина қаласында 4 күн мешіт жанында, Мекке қаласында 7 күн әл-Харамға жақын 5★ қонақүй).
- Топтамаға кіретін 9 негізгі қызмет:
  1. Әуе билеті (Шымкент - Мәдина / Жидда - Шымкент екі жаққа)
  2. Ресми Ұмра Визасы мен толық медициналық сақтандыру
  3. 5★ Премиум Қонақүйлер (Мекке мен Мәдинада мешіт жанындағы жайлы бөлмелер)
  4. VIP Трансфер (жайлы, салқындатқышы бар люкс автобустармен қатынау)
  5. Ұстаз жетекшілігі (Әлфараби ұстазбен амалдар, дұғалар, күнделікті уағыздар)
  6. Тарихи Зиярат (Ұхыд тауы, Құба мешіті, Қос құбыла, Нұр тауы, Сауыр үңгірі)
  7. Толық Тамақтану (3-4 мезгіл швед үстелі және халал асхана)
  8. 2x23 кг Багаж + 7 кг қол жүгі + 5 литр Зәмзәм сыйы
  9. Қажылар жинағы (ихрам, сөмке, бейдж, жолбасшы кітапша)
- Құжаттар: Жарамдылық мерзімі кемінде 6 ай қалған шетелдік төлқұжат және 3x4 көлеміндегі фотосурет (визаны толық агенттік рәсімдейді).
- Дайындық: Сапар алдында қажыларға 3 апталық тегін рухани және практикалық дәрістер өткізіледі.
- Ерекшелігі: Әр қажыға жеке жанашырлық пен қамқорлық, жан тыныштығы.

Жауап беру ережелерің:
1. Тек қазақ тілінде, өте сыпайы, жылы, сенімді әрі сауатты жауап бер.
2. Мәліметті жинақы, түсінікті (қажет болса қысқа тізіммен) жеткіз.
3. Егер қолданушы нақты бағаларды, жақын күндерді немесе орын брондауды сұраса, сұрағына қысқа мәлімет беріп: «Толық кеңес алу немесе орын брондау үшін төмендегі WhatsApp батырмасы арқылы менеджерімізбен байланысыңыз» деп бағытта.
"""

def record_lead_to_excel(name, phone, city, note="", source="Сайт"):
    """
    Thread-safe function to record incoming leads into:
    1. leads.xlsx (Native Microsoft Excel file)
    2. leads.csv  (UTF-8 with BOM for 1-click Excel opening)
    3. leads.json (JSON database backup)
    """
    with lead_lock:
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 1. Save to Excel (.xlsx)
        try:
            if not os.path.exists(EXCEL_FILE):
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Өтінімдер (Лидтер)"
                headers = ["№", "Күні мен уақыты", "Аты-жөні", "Телефон нөмірі", "Қаласы", "Қызықтыратын тур / Сұрағы", "Қайдан түсті"]
                ws.append(headers)
                
                header_fill = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                wb = openpyxl.load_workbook(EXCEL_FILE)
                ws = wb.active

            row_num = ws.max_row
            row_data = [row_num, now_str, name, phone, city, note, source]
            ws.append(row_data)

            # Style data row
            border_thin = Border(
                left=Side(style='thin', color='CCCCCC'),
                right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'),
                bottom=Side(style='thin', color='CCCCCC')
            )
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num + 1, column=col_idx)
                cell.border = border_thin
                cell.alignment = Alignment(vertical="center")

            # Adjust column widths
            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

            wb.save(EXCEL_FILE)
        except Exception as e:
            print(f"[ERROR] Failed to save Excel file: {e}")

        # 2. Save to CSV (UTF-8 BOM for Kazakh letters in Excel)
        try:
            write_header = not os.path.exists(CSV_FILE)
            with open(CSV_FILE, "a", encoding="utf-8-sig") as f:
                if write_header:
                    f.write("№,Күні мен уақыты,Аты-жөні,Телефон нөмірі,Қаласы,Қызықтыратын тур,Қайдан түсті\n")
                f.write(f'"{row_num}","{now_str}","{name}","{phone}","{city}","{note}","{source}"\n')
        except Exception as e:
            print(f"[ERROR] Failed to save CSV file: {e}")

        # 3. Save to JSON database
        try:
            leads = []
            if os.path.exists(JSON_FILE):
                try:
                    with open(JSON_FILE, "r", encoding="utf-8") as f:
                        leads = json.load(f)
                except Exception:
                    leads = []
            leads.append({
                "id": row_num,
                "timestamp": now_str,
                "name": name,
                "phone": phone,
                "city": city,
                "note": note,
                "source": source
            })
            with open(JSON_FILE, "w", encoding="utf-8") as f:
                json.dump(leads, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"[ERROR] Failed to save JSON database: {e}")

        print(f"[LEAD RECORDED] №{row_num} - {name} ({phone}, {city}) -> {EXCEL_FILE}")
        return row_num


class RangeRequestHandler(SimpleHTTPRequestHandler):
    """
    HTTP server with support for:
    1. HTTP 206 Partial Content (Range requests) for lag-free video seeking.
    2. POST /api/lead to record user data to Excel, CSV, and JSON database.
    3. GET /api/leads/download to download the Excel database anytime.
    """
    def end_headers(self):
        self.send_header('Accept-Ranges', 'bytes')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')
        self.end_headers()

    def do_POST(self):
        if self.path == '/api/chat':
            content_length = int(self.headers.get('Content-Length', 0))
            post_body = self.rfile.read(content_length).decode('utf-8')
            
            try:
                data = json.loads(post_body)
            except Exception:
                self.send_response(400)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "error", "message": "Invalid JSON"}).encode('utf-8'))
                return

            question = data.get('question', '').strip()
            if not question:
                self.send_response(400)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "error", "message": "Сұрақты енгізіңіз"}).encode('utf-8'))
                return

            try:
                req_payload = {
                    "model": "alemllm",
                    "messages": [
                        {"role": "system", "content": FARAB_SYSTEM_PROMPT},
                        {"role": "user", "content": question}
                    ]
                }
                api_req = urllib.request.Request(
                    ALEM_API_URL,
                    data=json.dumps(req_payload).encode('utf-8'),
                    headers={
                        'Content-Type': 'application/json',
                        'Authorization': f'Bearer {ALEM_API_KEY}'
                    }
                )
                with urllib.request.urlopen(api_req, timeout=25) as resp:
                    api_res = json.loads(resp.read().decode('utf-8'))
                    answer = api_res.get('choices', [{}])[0].get('message', {}).get('content', '')

                # Build prefilled WhatsApp escalation URL
                wa_question_text = f"Ассалаумағалейкум! Farab Tour AI кеңесшісінде сұрақ қойдым:\n\n«{question}»\n\nОсы сауал бойынша толық кеңес алып, орын брондағым келеді."
                wa_url = f"https://wa.me/77474983298?text={urllib.parse.quote(wa_question_text)}"

                self.send_response(200)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({
                    "status": "success",
                    "answer": answer,
                    "question": question,
                    "whatsapp_url": wa_url
                }, ensure_ascii=False).encode('utf-8'))
            except Exception as e:
                print(f"[AI CHAT ERROR] {e}")
                self.send_response(500)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.send_header('Access-Control-Allow-Origin', '*')
                self.end_headers()
                self.wfile.write(json.dumps({
                    "status": "error",
                    "message": "AI жауабын алу кезінде қате орын алды. WhatsApp арқылы тікелей сұрай аласыз.",
                    "whatsapp_url": f"https://wa.me/77474983298?text={urllib.parse.quote('Ассалаумағалейкум! Farab Tour Ұмра сапары бойынша кеңес алғым келеді.')}"
                }, ensure_ascii=False).encode('utf-8'))
            return

        elif self.path == '/api/lead':
            content_length = int(self.headers.get('Content-Length', 0))
            post_body = self.rfile.read(content_length).decode('utf-8')
            
            try:
                data = json.loads(post_body)
            except Exception:
                self.send_response(400)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "error", "message": "Invalid JSON"}).encode('utf-8'))
                return

            name = data.get('name', '').strip()
            phone = data.get('phone', '').strip()
            city = data.get('city', '').strip() or 'Көрсетілмеген'
            note = data.get('note', '').strip()
            source = data.get('source', 'Сайт')

            if not name or not phone:
                self.send_response(400)
                self.send_header('Content-type', 'application/json; charset=utf-8')
                self.end_headers()
                self.wfile.write(json.dumps({"status": "error", "message": "Аты-жөні мен телефон нөмірін толтырыңыз"}).encode('utf-8'))
                return

            # Save lead to Excel and database
            lead_id = record_lead_to_excel(name, phone, city, note, source)

            # Build WhatsApp message text
            wa_text = f"Сәлеметсіз бе! Мен Farab Tour сайтынан өтінім қалдырдым:\n\n"
            wa_text += f"👤 Аты-жөнім: {name}\n"
            wa_text += f"📱 Телефон: {phone}\n"
            wa_text += f"📍 Қала: {city}\n"
            if note:
                wa_text += f"🕋 Сұрағым/Тур: {note}\n"
            wa_text += f"\nТолық кеңес алуға дайынмын!"

            response_data = {
                "status": "success",
                "lead_id": lead_id,
                "message": "Өтініміңіз базаға сәтті сақталды!",
                "whatsapp_text": wa_text,
                "whatsapp_url": f"https://wa.me/77474983298?text={json.dumps(wa_text)[1:-1]}"
            }

            self.send_response(200)
            self.send_header('Content-type', 'application/json; charset=utf-8')
            self.send_header('Access-Control-Allow-Origin', '*')
            self.end_headers()
            self.wfile.write(json.dumps(response_data, ensure_ascii=False).encode('utf-8'))
        else:
            self.send_error(404, "Endpoint not found")

    def send_head(self):
        # Handle Excel download endpoint
        if self.path == '/api/leads/download':
            if not os.path.exists(EXCEL_FILE):
                self.send_error(404, "Әзірге өтінімдер тіркелмеген")
                return None
            try:
                f = open(EXCEL_FILE, 'rb')
            except OSError:
                self.send_error(404, "Excel файлын ашу мүмкін болмады")
                return None
            fs = os.fstat(f.fileno())
            self.send_response(200)
            self.send_header("Content-type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", 'attachment; filename="farab_tour_leads.xlsx"')
            self.send_header("Content-Length", str(fs[6]))
            self.end_headers()
            return f

        if 'Range' not in self.headers:
            self.range = None
            return super().send_head()
        
        path = self.translate_path(self.path)
        if not os.path.exists(path) or os.path.isdir(path):
            return super().send_head()
            
        ctype = self.guess_type(path)
        try:
            f = open(path, 'rb')
        except OSError:
            self.send_error(404, "File not found")
            return None

        fs = os.fstat(f.fileno())
        total_len = fs[6]
        
        range_header = self.headers.get('Range', '')
        if not range_header.startswith('bytes='):
            return super().send_head()
            
        ranges = range_header[6:].split('-')
        start = int(ranges[0]) if ranges[0] else 0
        end = int(ranges[1]) if len(ranges) > 1 and ranges[1] else total_len - 1
        
        if start >= total_len or end >= total_len or start > end:
            self.send_error(416, "Requested Range Not Satisfiable")
            return None
            
        self.send_response(206)
        self.send_header("Content-type", ctype)
        self.send_header("Content-Range", f"bytes {start}-{end}/{total_len}")
        self.send_header("Content-Length", str(end - start + 1))
        self.send_header("Accept-Ranges", "bytes")
        self.send_header("Last-Modified", self.date_time_string(fs.st_mtime))
        self.end_headers()
        
        f.seek(start)
        self.range = (start, end)
        return f

    def copyfile(self, source, outputfile):
        if not hasattr(self, 'range') or not self.range:
            return super().copyfile(source, outputfile)
        start, end = self.range
        bytes_to_send = end - start + 1
        buffer_size = 64 * 1024
        while bytes_to_send > 0:
            read_len = min(buffer_size, bytes_to_send)
            buf = source.read(read_len)
            if not buf:
                break
            try:
                outputfile.write(buf)
            except (ConnectionResetError, BrokenPipeError):
                break
            bytes_to_send -= len(buf)

def run():
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 8080
    server_address = ('0.0.0.0', port)
    httpd = HTTPServer(server_address, RangeRequestHandler)
    print(f"Farab Tour Server (HTTP 206 Range + Excel Lead API) running on port {port}...")
    httpd.serve_forever()

if __name__ == '__main__':
    run()
